from datetime import datetime
from dotenv import load_dotenv

#from app.api.v1 import RESUME_VERSIONS, validate_text_soft
#from app.recruiter.ranking_engine import rank_resumes_against_jd
load_dotenv()
from fastapi import FastAPI, UploadFile, File, HTTPException, Body, Form
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import os
import shutil
import uuid, datetime
from typing import Optional, List
from io import BytesIO
from app.parsing.resume_parser import parse_resume
from app.parsing.jd_parser import parse_jd
from app.utils.text_validator import validate_min_words
from app.skills.skill_extractor import extract_skills
from app.matching.semantic_matcher import semantic_match
from app.matching.jd_skill_classifier import classify_jd_skills
from app.matching.skill_matcher import match_skills
from app.matching.weighted_scorer import calculate_ats_score
from app.matching.explainability import explain_score
from app.analysis.experience_analyzer import calculate_experience_score,extract_experience_years
from app.analysis.project_analyzer import calculate_project_relevance_score
from app.analysis.ats_checker import calculate_ats_format_score
#from app.recommendations.engine import generate_skill_gap_recommendations
from app.analysis.final_scorer import calculate_final_ats_score
from app.analysis.project_analyzer import analyze_projects_core
from app.matching.resume_jd_matcher import calculate_resume_match_score_full,calculate_resume_jd_similarity
from app.role_intelligence.role_detector import detect_role, calculate_role_relevance_score
from app.ai_engine.ai_guard import ai_safe_execute
#from app.analysis.fraud_detector import detect_resume_fraud
from app.role_intelligence.role_detector import (
    detect_role,
    calculate_role_relevance_score,
)

from app.ai_engine.skill_fallback import (
    detect_unknown_skills,
    classify_unknown_skills,
)

from app.analysis.section_analyzer import generate_section_feedback
from app.ai_engine.section_feedback_ai import generate_section_feedback_ai
from app.ai_engine.resume_rewrite_ai import rewrite_resume_for_jd
#from app.ai_engine.prompts import INDIVIDUAL_BULLET_IMPROVEMENT_PROMPT, SUMMARY_IMPROVEMENT_PROMPT
#from app.recommendations.engine import generate_skill_gap_recommendations





# ✅ ONE app only
app = FastAPI(title="AI Resume Analyzer Backend")

# ✅ CORS Configuration - CRITICAL for Frontend Integration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allow all origins (or specify frontend URL like ["http://localhost:3000"])
    allow_credentials=True,
    allow_methods=["*"],  # Allow all HTTP methods (GET, POST, PUT, DELETE, etc.)
    allow_headers=["*"],  # Allow all headers
)

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)


# -----------------------------
# Health Check
# -----------------------------
@app.get("/")
def health_check():
    return {"status": "Backend is running"}

"""
# -----------------------------
# Upload Resume
# -----------------------------
@app.post("/upload-resume")
async def upload_resume(file: UploadFile = File(...)):
    if not file.filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Unsupported resume format")

    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    text = parse_resume(file_path)

    return {
        "filename": file.filename,
        "extracted_text_preview": text[:1000]
    }


# -----------------------------
# Upload JD
# -----------------------------
@app.post("/upload-jd")
async def upload_jd(file: UploadFile = File(...)):
    if not file.filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Unsupported JD format")

    file_path = os.path.join(UPLOAD_DIR, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    text = parse_jd(file_path)

    return {
        "filename": file.filename,
        "extracted_text_preview": text[:1000]
    }


# -----------------------------
# Extract Skills (raw text)
# -----------------------------
@app.post("/extract-skills")
async def extract_skills_api(text: str = Body(..., media_type="text/plain")):
    skills = extract_skills(text)
    return {
        "extracted_skills": skills,
        "count": len(skills)
    }

"""
# -----------------------------
#  ANALYZE ENDPOINT
# -----------------------------

@app.post("/analyze")
async def analyze(
    resume_file: Optional[UploadFile] = File(None),
    jd_file: Optional[UploadFile] = File(None),
    resume_text: Optional[str] = Form(None),
    jd_text: Optional[str] = Form(None),
):
    
# ---------------------------
# Resume input handling
# ---------------------------
    if resume_file and resume_file.filename:
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)
        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())
        resume_text = parse_resume(resume_path)

    elif resume_text and resume_text.strip() and resume_text != "string":
        resume_text = resume_text.strip()

    else:
        resume_text = ""

# ---------------------------
# JD input handling
# ---------------------------
    if jd_file and jd_file.filename:
        jd_path = os.path.join(UPLOAD_DIR, jd_file.filename)
        with open(jd_path, "wb") as f:
            f.write(await jd_file.read())
        jd_text = parse_jd(jd_path)

    elif jd_text and jd_text.strip() and jd_text != "string":
        jd_text = jd_text.strip()

    else:
        jd_text = ""

    if not resume_text or not jd_text:
        raise HTTPException(
            status_code=400,
            detail="Resume or Job Description content could not be extracted"
        )

# ---------------------------
# Validate minimum word count
# ---------------------------
    resume_valid = validate_min_words(resume_text)
    jd_valid = validate_min_words(jd_text)

    if not resume_valid and not jd_valid:
        raise HTTPException(
            status_code=400,
            detail="Resume and Job Description must each contain at least 20 words"
        )

    if not resume_valid:
        raise HTTPException(
            status_code=400,
            detail="Resume must contain at least 20 words"
        )

    if not jd_valid:
        raise HTTPException(
            status_code=400,
            detail="Job Description must contain at least 20 words"
        )


    


    # ---------- 3. Extract skills ----------
    resume_skills = extract_skills(resume_text)
    jd_skills = extract_skills(jd_text)

        # ---------- AI Fallback Skill Detection ----------
    unknown_resume_skills = detect_unknown_skills(resume_skills, jd_skills)
    ai_classified_skills = classify_unknown_skills(unknown_resume_skills)

    # Merge AI-detected skills
    # ai_classified_skills is a LIST of dicts
    # -------- Merge AI-detected skills (SAFE) --------
    for item in ai_classified_skills:
        # Case 1: dict
        if isinstance(item, dict):
            skill_name = item.get("skill")

        # Case 2: string
        elif isinstance(item, str):
            skill_name = item

        else:
            continue

        if skill_name and skill_name not in resume_skills:
            resume_skills.append(skill_name)


    # ---------- END AI Fallback Skill Detection ----------
    if not resume_skills and not jd_skills:
        return {
            "error": "No skills detected in resume and job description",
            "resume_skills": [],
            "jd_skills": [],
            "ats_score": 0,
            "score_explanation": [
                "Neither resume nor job description contains recognizable skills"
            ]
        }

    if not resume_skills:
        return {
            "error": "No skills detected in resume",
            "resume_skills": [],
            "jd_skills": jd_skills,
            "ats_score": 0,
            "score_explanation": [
                "Resume does not contain recognizable skills"
            ]
        }

    if not jd_skills:
            return {
                "error": "No skills detected in job description",
                "resume_skills": resume_skills,
                "jd_skills": [],
                "ats_score": 0,
                "score_explanation": [
                    "Job description does not contain recognizable skills"
                ]
        }
    

    # ---------- 4. Rule-based matching ----------
    rule_matches = set(resume_skills) & set(jd_skills)

    # ---------- 5. Semantic matching ----------
    semantic_result = semantic_match(resume_skills, jd_skills)
    semantic_matches = set(semantic_result["matched_skills"])

    # ---------- 6. Merge results (DAY 8 CORE) ----------
    final_matched = sorted(rule_matches | semantic_matches)
    final_missing = sorted(set(jd_skills) - set(final_matched))

    #---skill gap recommendations---    
    #skill_gap_recommendations = generate_skill_gap_recommendations(final_missing)

    # ---------------------------
    # Skill Score (JD Coverage)
    # ---------------------------
    if jd_skills:
        skill_score = round((len(final_matched) / len(jd_skills)) * 100)
    else:
        skill_score = 0

    

    #---------experience score------------
    experience_result = calculate_experience_score(resume_text, jd_text)
    experience_score = experience_result.get("experience_score", 0)

    #---------project relevance score------------------------------

    project_result = analyze_projects_core(
        resume_text=resume_text,
        resume_skills=resume_skills,
        jd_skills=jd_skills,
)
    project_score = project_result["project_score"]


    #---------ATS format score----------------
    ats_format_result = calculate_ats_format_score(resume_text)
    ats_format_score = ats_format_result["ats_format_score"]

    #---------role relevance score----------------
        # Detect role from JD
    detected_role = detect_role(jd_skills)


    # Calculate role relevance score
    role_result = calculate_role_relevance_score(
        resume_skills=resume_skills,
        detected_role=detected_role
    )

    role_score = role_result["role_score"]



    # ---------- 7. Classify JD skills ----------
    jd_classification = classify_jd_skills(jd_text, jd_skills)

    # -------- Full Semantic Similarity --------
    semantic_result = calculate_resume_jd_similarity(
        resume_text=resume_text,
        jd_text=jd_text
    )

    semantic_score = semantic_result["resume_match_score"]

   # -------- Full Resume Match Score --------
    full_match_result = calculate_resume_match_score_full(
        semantic_score=semantic_score,
        skill_score=skill_score,
        experience_score=experience_score,
        project_score=project_score,
        role_score=role_score,
    )

    resume_match_score = full_match_result["resume_match_score"]

    # ---------- 8. Weighted ATS score ----------
    ats_score = calculate_ats_score(
        final_matched,
        jd_classification["must_have"],
        jd_classification["good_to_have"]
    )

    # ---------- 9. Explainability ----------
    score_explanation = explain_score(
        final_matched,
        jd_classification["must_have"],
        jd_classification["good_to_have"]
    )

    # ---------- Edge case: zero ATS score ----------
    if ats_score == 0:
        score_explanation.append("No must-have skills matched")

    # -----------------------------
    # SECTION-WISE FEEDBACK (AI)
    # -----------------------------

    section_feedback_raw= ai_safe_execute(
        ai_function=generate_section_feedback_ai,
        fallback_function=lambda resume_text, jd_text: generate_section_feedback(resume_text, jd_skills),
        resume_text=resume_text,
        jd_text=jd_text
    )

    section_feedback = section_feedback_raw["result"]

     # ---------- AI reusme rewrite ----------
    #ai_improved = rewrite_resume_for_jd(resume_text, jd_text)


    #--fraud detection--
    #years_of_experience = extract_experience_years(resume_text)
    #fraud_report = detect_resume_fraud(resume_skills, years_of_experience)

    final_ats = calculate_final_ats_score(
        skill_score=resume_match_score,
        experience_score=experience_score,
        project_score=project_score,
        ats_format_score=ats_format_score,
        role_score=role_score,
    )

    print("USING FINAL ATS:", final_ats)



    return {
        #"resume_skills": resume_skills,
        #"jd_skills": jd_skills,
        "ats_score" : final_ats["ats_score"] ,
        "resume_match_score": resume_match_score,
        "role": detected_role,
        "matched_skills": final_matched,
        "missing_skills": final_missing,
        "must_have_skills": jd_classification["must_have"],
        "good_to_have_skills": jd_classification["good_to_have"],
        "skill_coverage": {
        "matched_percentage": skill_score,
        "missing_percentage": 100 - skill_score
    },
    "chart_data": {
        "resume_skill_count": len(resume_skills),
        "jd_skill_count": len(jd_skills),
    },
     "ai_insights": section_feedback,
    }

        # "skill_gap_recommendations": skill_gap_recommendations,
        #"ai_resume_improvements": ai_improved,
        #"score_explanation": score_explanation,
        #"fraud_report": fraud_report,
        #"_debug": {
        #    "rule_matches_count": len(rule_matches),
        #    "semantic_matches_count": len(semantic_matches)
        #}





# -----------------------------
# Resume Rewrite API
# -----------------------------
@app.post("/rewrite-resume")
async def rewrite_resume(
    resume_file: Optional[UploadFile] = File(None),
    jd_file: Optional[UploadFile] = File(None),
    resume_text: Optional[str] = Form(None),
    jd_text: Optional[str] = Form(None)
):
    """
    AI-powered resume rewrite endpoint for 3rd page UI.
    
    Returns improved resume sections:
    - Summary: Optimized professional summary
    - Experience: List of improved experience bullets
    - Projects: List of improved project descriptions
    
    Two input modes:
    1. Frontend integration: Send resume_text + jd_text directly
    2. Swagger testing: Upload resume_file + jd_file
    """
    
    # ---------------------------
    # Input Handling
    # ---------------------------
    if resume_file and resume_file.filename:
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)
        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())
        resume_text = parse_resume(resume_path)
    
    if jd_file and jd_file.filename:
        jd_path = os.path.join(UPLOAD_DIR, jd_file.filename)
        with open(jd_path, "wb") as f:
            f.write(await jd_file.read())
        jd_text = parse_jd(jd_path)
    
    if not resume_text or not jd_text:
        raise HTTPException(
            status_code=400,
            detail="Provide either files or text for both resume and job description"
        )

    # ---------------------------
    # Call AI Rewrite Function
    # ---------------------------
    try:
        rewrite_result = rewrite_resume_for_jd(resume_text, jd_text)
        
        # Check if AI returned an error
        if "error" in rewrite_result:
            return {
                "status": "error",
                "message": rewrite_result["error"],
                "rewritten_resume": {
                    "summary": "",
                    "experience": [],
                    "projects": []
                },
                "metadata": {
                    "model": "gemini-2.5-flash",
                    "timestamp": datetime.datetime.now().isoformat(),
                    "ai_used": False
                }
            }
        
        # Success - return rewritten sections
        return {
            "status": "success",
            "message": "Resume successfully rewritten for target job description",
            "rewritten_resume": {
                "summary": rewrite_result.get("summary", ""),
                "experience": rewrite_result.get("experience", []),
                "projects": rewrite_result.get("projects", [])
            },
            "metadata": {
                "model": "gemini-2.5-flash",
                "timestamp": datetime.datetime.now().isoformat(),
                "ai_used": True,
                "experience_count": len(rewrite_result.get("experience", [])),
                "project_count": len(rewrite_result.get("projects", []))
            }
        }
        
    except Exception as e:
        # Handle any errors gracefully
        return {
            "status": "error",
            "message": f"Rewrite failed: {str(e)}",
            "rewritten_resume": {
                "summary": "",
                "experience": [],
                "projects": []
            },
            "metadata": {
                "model": "gemini-2.5-flash",
                "timestamp": datetime.datetime.now().isoformat(),
                "ai_used": False
            }
        }



# ----------------------------------
# Download Optimized Resume
# ----------------------------------
@app.post("/download-resume")
async def download_resume(request_body: dict = Body(...)):
    """
    Download professional resume with improvements applied.
    
    Accepts output directly from /rewrite-resume endpoint.
    
    Input: JSON with summary, experience, projects, format
    Output: Downloadable file (txt, docx, or json)
    """
    
    try:
        # Handle both direct format and wrapped rewritten_resume format
        if "rewritten_resume" in request_body:
            data = request_body.get("rewritten_resume", {})
        else:
            data = request_body
        
        summary = data.get("summary", "")
        experience_input = data.get("experience", [])
        projects_input = data.get("projects", [])
        file_format = data.get("format", "txt").lower()
        
        # Clean summary
        summary = str(summary).strip() if summary else ""
        
        # Process experience list
        experience = []
        if experience_input:
            for exp in experience_input:
                if isinstance(exp, dict):
                    # Handle {title, company, date, bullets} format
                    title = exp.get("title", "")
                    company = exp.get("company", "")
                    date = exp.get("date", "")
                    bullets = exp.get("bullets", [])
                    
                    # Build experience entry
                    exp_entry = ""
                    if title:
                        exp_entry = title
                    if company:
                        exp_entry += f" at {company}" if exp_entry else company
                    if date:
                        exp_entry += f" ({date})" if exp_entry else date
                    
                    if exp_entry:
                        experience.append(exp_entry)
                    
                    # Add bullets
                    if isinstance(bullets, list):
                        for bullet in bullets:
                            if bullet and str(bullet).strip():
                                experience.append(f"  • {str(bullet).strip()}")
                
                elif isinstance(exp, str):
                    if exp.strip():
                        experience.append(exp.strip())
        
        # Process projects list
        projects = []
        if projects_input:
            for proj in projects_input:
                if isinstance(proj, dict):
                    # Handle {title, bullets, description, model, technologies} format
                    title = proj.get("title", "")
                    description = proj.get("description", "")
                    bullets = proj.get("bullets", [])
                    model = proj.get("model", "")
                    technologies = proj.get("technologies", "")
                    
                    # Build project entry
                    proj_entry = ""
                    if title:
                        proj_entry = title
                    if description:
                        proj_entry += f": {description}" if proj_entry else description
                    
                    if proj_entry:
                        projects.append(proj_entry)
                    
                    # Add bullets if they exist
                    if isinstance(bullets, list):
                        for bullet in bullets:
                            if bullet and str(bullet).strip():
                                projects.append(f"  • {str(bullet).strip()}")
                    
                    # Add model info if exists
                    if model:
                        projects.append(f"  • Model: {model}")
                    if technologies:
                        projects.append(f"  • Technologies: {technologies}")
                
                elif isinstance(proj, str):
                    if proj.strip():
                        projects.append(proj.strip())
        
        # Build resume content
        resume_content = ""
        
        # Add Summary
        if summary:
            resume_content += "PROFESSIONAL SUMMARY\n"
            resume_content += "=" * 70 + "\n"
            resume_content += summary + "\n\n"
        
        # Add Experience
        if experience:
            resume_content += "PROFESSIONAL EXPERIENCE\n"
            resume_content += "=" * 70 + "\n"
            for item in experience:
                if item.startswith("  •"):
                    resume_content += f"{item}\n"
                else:
                    resume_content += f"{item}\n"
            resume_content += "\n"
        
        # Add Projects
        if projects:
            resume_content += "PROJECTS\n"
            resume_content += "=" * 70 + "\n"
            for item in projects:
                if item.startswith("  •"):
                    resume_content += f"{item}\n"
                else:
                    resume_content += f"{item}\n"
            resume_content += "\n"
        
        # Return based on format
        if file_format == "txt":
            file_bytes = resume_content.encode('utf-8')
            return StreamingResponse(
                iter([file_bytes]),
                media_type="text/plain",
                headers={
                    "Content-Disposition": "attachment; filename=resume_optimized.txt",
                    "Content-Length": str(len(file_bytes))
                }
            )
        
        elif file_format == "json":
            import json
            resume_json = {
                "summary": summary,
                "experience": experience,
                "projects": projects,
                "generated_at": datetime.datetime.now().isoformat()
            }
            file_bytes = json.dumps(resume_json, indent=2).encode('utf-8')
            return StreamingResponse(
                iter([file_bytes]),
                media_type="application/json",
                headers={
                    "Content-Disposition": "attachment; filename=resume_optimized.json",
                    "Content-Length": str(len(file_bytes))
                }
            )
        
        elif file_format == "docx":
            try:
                from docx import Document
                from docx.shared import Pt
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                
                doc = Document()
                
                if summary:
                    heading = doc.add_heading("PROFESSIONAL SUMMARY", level=1)
                    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    doc.add_paragraph(summary)
                    doc.add_paragraph()
                
                if experience:
                    heading = doc.add_heading("PROFESSIONAL EXPERIENCE", level=1)
                    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for item in experience:
                        if item.startswith("  •"):
                            doc.add_paragraph(item.replace("  • ", ""), style='List Bullet')
                        else:
                            doc.add_paragraph(item)
                    doc.add_paragraph()
                
                if projects:
                    heading = doc.add_heading("PROJECTS", level=1)
                    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for item in projects:
                        if item.startswith("  •"):
                            doc.add_paragraph(item.replace("  • ", ""), style='List Bullet')
                        else:
                            doc.add_paragraph(item)
                
                file_stream = BytesIO()
                doc.save(file_stream)
                file_bytes = file_stream.getvalue()
                
                return StreamingResponse(
                    iter([file_bytes]),
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={
                        "Content-Disposition": "attachment; filename=resume_optimized.docx",
                        "Content-Length": str(len(file_bytes))
                    }
                )
            except ImportError:
                return {
                    "status": "error",
                    "message": "DOCX format requires python-docx. Use txt or json instead."
                }
        
        else:
            return {
                "status": "error",
                "message": "Format must be txt, json, or docx"
            }
    
    except Exception as e:
        return {
            "status": "error",
            "message": f"Error generating resume: {str(e)}"
        }


# INTERNAL DEV ENDPOINTS (Not used by frontend)

# -----------------------------
# Experience Analysis API
# -----------------------------
@app.post("/analyze-experience")
async def analyze_experience(
        resume_file: UploadFile = File(...),
        jd_file: UploadFile = File(...)
    ):
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)
        jd_path = os.path.join(UPLOAD_DIR, jd_file.filename)

        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())

        with open(jd_path, "wb") as f:
            f.write(await jd_file.read())

        resume_text = parse_resume(resume_path)
        jd_text = parse_jd(jd_path)

        experience_result = calculate_experience_score(resume_text, jd_text)

        return {
            "experience_analysis": experience_result
        }


    # -----------------------------
    # Project Relevance Analysis API
    # -----------------------------
@app.post("/analyze-projects")
async def analyze_projects_api(
        resume_file: UploadFile = File(...),
        jd_file: UploadFile = File(...),
    ):
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)
        jd_path = os.path.join(UPLOAD_DIR, jd_file.filename)

        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())

        with open(jd_path, "wb") as f:
            f.write(await jd_file.read())

        resume_text = parse_resume(resume_path)
        jd_text = parse_jd(jd_path)

        resume_skills = extract_skills(resume_text)
        jd_skills = extract_skills(jd_text)

        project_result = calculate_project_relevance_score(
            resume_text=resume_text,
            jd_skills=jd_skills,
            known_skills=resume_skills,
        )

        return {
            "project_analysis": project_result
        }

    # -----------------------------
    # ATS Format Compatibility API
    # -----------------------------
@app.post("/analyze-ats-format")
async def analyze_ats_format(resume_file: UploadFile = File(...)):
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)

        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())

        resume_text = parse_resume(resume_path)

        ats_result = calculate_ats_format_score(resume_text)

        return {
            "ats_format_analysis": ats_result
        }

    # -----------------------------
    # Dynamic Skill Intelligence API
    # -----------------------------
@app.post("/analyze-dynamic-skills")
async def analyze_dynamic_skills(
        resume_file: UploadFile = File(...),
):
        resume_path = os.path.join(UPLOAD_DIR, resume_file.filename)

        with open(resume_path, "wb") as f:
            f.write(await resume_file.read())

        resume_text = parse_resume(resume_path)

        extracted_skills = extract_skills(resume_text)

        # Known skill taxonomy
        from app.skills.skill_list import SKILL_LIST

        # Detect unknown skills
        unknown_skills = detect_unknown_skills(
            extracted_skills=extracted_skills,
            known_skills=SKILL_LIST,
        )

        # Classify unknown skills using AI
        classified_skills = classify_unknown_skills(unknown_skills)

        return {
            "extracted_skills": extracted_skills,
            "unknown_skills": unknown_skills,
            "ai_detected_skills": classified_skills,
        }


"""
# ======================================================
# 2️⃣ MULTI-JD ANALYSIS (Resume vs many JDs)
# ======================================================
@router.post("/analyze/multi-jd")
async def analyze_multi_jd(
    resume_text: str = Form(...),
    jd_texts: List[str] = Form(...)
):
    validation_warnings = []
    validate_text_soft(resume_text, "Resume", validation_warnings)




    results = []
    for jd in jd_texts:
        validate_text_soft(jd_texts, "Job Description", validation_warnings)


        results.append(await analyze_resume(
            resume_text=resume_text,
            jd_text=jd
        ))


    return {"results": results}


# ======================================================
# 3️⃣ MULTI-RESUME RANKING (Recruiter)
# ======================================================@router.post("/rank-resumes")
async def rank_resumes(
    resumes: List[str] = Form(...),
    jd_text: str = Form(...)
):
    validation_warnings = []
    validate_text_soft(jd_text, "Job Description", validation_warnings)


    return rank_resumes_against_jd(resumes, jd_text)


# ======================================================
# 4️⃣ RESUME VERSION HISTORY
# ======================================================
@router.get("/resume/{resume_id}/versions")
def get_resume_versions(resume_id: str):
    return RESUME_VERSIONS.get(resume_id, [])
"""